from datetime import date
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from .models import Tarifa, Viaje, Grupo, Turista


def crear_grupo(viaje, *personas):
    grupo = Grupo.objects.create(viaje=viaje)
    for nombre, rappel, tirolesa in personas:
        Turista.objects.create(viaje=viaje, grupo=grupo, nombre_completo=nombre,
                               hace_rappel=rappel, hace_tirolesa=tirolesa)
    return grupo


class PreciosTests(TestCase):
    def setUp(self):
        self.viaje = Viaje.objects.create(fecha_salida=date(2026, 11, 1))

    def saldo(self, nombre):
        return Turista.objects.get(nombre_completo=nombre).saldo

    def test_viaje_copia_la_tarifa_vigente(self):
        self.assertEqual(self.viaje.precio_paquete, Decimal('270'))
        self.assertEqual(self.viaje.precio_rappel, Decimal('120'))
        self.assertEqual(self.viaje.descuento_grupo, Decimal('20'))

    def test_persona_sola_sin_descuento(self):
        crear_grupo(self.viaje, ('Ana', True, True))
        # 270 + 120 + 120 - 50 de reserva
        self.assertEqual(self.saldo('Ana'), Decimal('460'))

    def test_grupo_de_dos_sin_descuento(self):
        crear_grupo(self.viaje, ('Ana', False, False), ('Beto', False, False))
        self.assertEqual(self.saldo('Ana'), Decimal('220'))
        self.assertEqual(self.saldo('Beto'), Decimal('220'))

    def test_grupo_de_tres_descuento_por_persona_solo_al_paquete(self):
        crear_grupo(self.viaje, ('Ana', False, False), ('Beto', True, False), ('Caro', True, True))
        # Todos reciben 20 de descuento, incluidos los dos primeros que se guardaron antes
        self.assertEqual(self.saldo('Ana'), Decimal('200'))
        self.assertEqual(self.saldo('Beto'), Decimal('320'))
        self.assertEqual(self.saldo('Caro'), Decimal('440'))

    def test_eliminar_del_grupo_quita_el_descuento(self):
        crear_grupo(self.viaje, ('Ana', False, False), ('Beto', False, False), ('Caro', False, False))
        Turista.objects.get(nombre_completo='Caro').delete()
        self.assertEqual(self.saldo('Ana'), Decimal('220'))

    def test_mover_a_otro_grupo_recalcula_ambos(self):
        g1 = crear_grupo(self.viaje, ('Ana', False, False), ('Beto', False, False), ('Caro', False, False))
        g2 = crear_grupo(self.viaje, ('Dani', False, False), ('Eva', False, False))
        caro = Turista.objects.get(nombre_completo='Caro')
        caro.grupo = g2
        caro.save()
        self.assertEqual(self.saldo('Ana'), Decimal('220'))
        self.assertEqual(self.saldo('Dani'), Decimal('200'))
        self.assertEqual(g1.turistas.count(), 2)

    def test_cambiar_tarifa_no_afecta_viajes_existentes(self):
        crear_grupo(self.viaje, ('Ana', False, False))
        tarifa = Tarifa.actual()
        tarifa.precio_paquete = Decimal('300')
        tarifa.save()
        Turista.objects.get(nombre_completo='Ana').save()
        self.assertEqual(self.saldo('Ana'), Decimal('220'))
        nuevo = Viaje.objects.create(fecha_salida=date(2026, 12, 1))
        self.assertEqual(nuevo.precio_paquete, Decimal('300'))

    def test_cambiar_precio_del_viaje_recalcula_saldos(self):
        crear_grupo(self.viaje, ('Ana', True, False))
        self.viaje.precio_rappel = Decimal('150')
        self.viaje.save()
        self.assertEqual(self.saldo('Ana'), Decimal('370'))


class ApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.viaje = Viaje.objects.create(fecha_salida=date(2026, 11, 1))

    def test_crear_grupo_desde_un_mensaje(self):
        datos = {'viaje': self.viaje.id, 'turistas': [
            {'nombre_completo': 'ana pérez', 'hace_rappel': True},
            {'nombre_completo': 'beto', 'hace_tirolesa': True},
            {'nombre_completo': 'caro'},
        ]}
        r = self.client.post('/api/grupos/', datos, format='json')
        self.assertEqual(r.status_code, 201, r.content)
        saldos = sorted(Decimal(t['saldo']) for t in r.json()['turistas'])
        self.assertEqual(saldos, [Decimal('200'), Decimal('320'), Decimal('320')])
        self.assertEqual(r.json()['turistas'][0]['nombre_completo'], 'Ana Pérez')

    def test_saldo_no_se_puede_escribir_desde_la_api(self):
        r = self.client.post('/api/turistas/', {'viaje': self.viaje.id, 'nombre_completo': 'X', 'saldo': 1}, format='json')
        self.assertEqual(Decimal(r.json()['saldo']), Decimal('220'))

    def test_tarifa_get_y_put(self):
        r = self.client.put('/api/tarifa/', {'precio_rappel': '130'}, format='json')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(Decimal(self.client.get('/api/tarifa/').json()['precio_rappel']), Decimal('130'))

    def test_excel_logistica_colorea_grupos(self):
        crear_grupo(self.viaje, ('Ana', False, False), ('Beto', False, False), ('Caro', False, False))
        crear_grupo(self.viaje, ('Dani', False, False))
        Turista.objects.create(viaje=self.viaje, nombre_completo='Suelto')
        r = self.client.get(f'/api/exportar-excel/{self.viaje.id}/logistica/')
        ws = load_workbook(BytesIO(r.content)).active
        filas = list(ws.iter_rows(min_row=5, values_only=True))
        self.assertEqual([f[2] for f in filas], ['G1', 'G1', 'G1', 'G2', '-', '5 pax'])
        self.assertNotEqual(ws['A5'].fill.fgColor.rgb, ws['A8'].fill.fgColor.rgb)
        self.assertEqual(ws['A5'].fill.fgColor.rgb, ws['I7'].fill.fgColor.rgb)
