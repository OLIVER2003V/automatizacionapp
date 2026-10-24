import os
import json
from google import genai

from rest_framework import viewsets, mixins, status
from rest_framework.views import APIView
from rest_framework.response import Response

from django.db.models import F
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

from .models import Tarifa, Viaje, Grupo, Turista
from .serializers import TarifaSerializer, ViajeSerializer, GrupoSerializer, TuristaSerializer

class ViajeViewSet(viewsets.ModelViewSet):
    queryset = Viaje.objects.all().order_by('-fecha_salida')
    serializer_class = ViajeSerializer

class TuristaViewSet(viewsets.ModelViewSet):
    queryset = Turista.objects.all().order_by('-fecha_registro')
    serializer_class = TuristaSerializer

class GrupoViewSet(mixins.CreateModelMixin, mixins.RetrieveModelMixin, mixins.DestroyModelMixin,
                   viewsets.GenericViewSet):
    queryset = Grupo.objects.all()
    serializer_class = GrupoSerializer

class TarifaView(APIView):
    """GET devuelve los precios vigentes; PUT/PATCH los cambia (solo afecta a viajes nuevos)."""
    def get(self, request):
        return Response(TarifaSerializer(Tarifa.actual()).data)

    def put(self, request):
        serializer = TarifaSerializer(Tarifa.actual(), data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(serializer.data)

    patch = put


# Ahora la API Key está oculta y segura. El cliente se crea al primer uso
# para que manage.py (migraciones, tests) funcione sin la clave.
_client = None

def obtener_cliente_ia():
    global _client
    if _client is None:
        _client = genai.Client(api_key=os.getenv('GEMINI_API_KEY'))
    return _client

class ProcesarMensajeIAView(APIView):
    def post(self, request):
        mensaje = request.data.get('mensaje', '')
        
        if not mensaje:
            return Response({"error": "No se envió ningún mensaje"}, status=status.HTTP_400_BAD_REQUEST)

        prompt = f"""
        Eres un analista de datos experto trabajando para una agencia de turismo en Bolivia. 
        Tu única tarea es extraer la información de los turistas de mensajes de WhatsApp desordenados con precisión absoluta.
        
        REGLAS LÓGICAS DE EXTRACCIÓN:
        1. CI (Cédula de Identidad): En Bolivia es un número de 6 a 8 dígitos. Suele estar cerca del nombre. Búscalo exhaustivamente.
        2. Celulares: Los celulares en Bolivia tienen EXACTAMENTE 8 dígitos y empiezan con 6 o 7. 
        3. Actividades: Si el turista dice "depende" para rappel o tirolesa, pon el valor en TRUE obligatoriamente.
        4. Observaciones Médicas: Si la persona NO TIENE alergias, NO TIENE fracturas o dice "ninguna", pon EXACTAMENTE null. ESTÁ PROHIBIDO escribir "No alergias", "Ninguna", o "No fracturas". Solo pon null. Si la actividad "depende", anota "Depende clima".

        REGLAS DE FORMATO ESTRICTAS:
        - Devuelve ÚNICAMENTE un arreglo (array) de objetos JSON.
        - NO incluyas etiquetas de markdown como ```json.

        ESTRUCTURA JSON REQUERIDA:
        [
            {{
                "nombre_completo": "Nombre",
                "ci": "Número o null",
                "edad": número o null,
                "celular": "Número o null",
                "contacto_referencia": "Número o null",
                "es_vegetariano": true/false,
                "condicion_medica": "Texto o null",
                "hace_rappel": true/false,
                "hace_tirolesa": true/false
            }}
        ]

        MENSAJE A PROCESAR:
        "{mensaje}"
        """

        try:
            response = obtener_cliente_ia().models.generate_content(
                model='gemini-2.5-flash',
                contents=prompt,
            )
            
            texto_limpio = response.text.replace('```json', '').replace('```', '').strip()
            datos_turistas = json.loads(texto_limpio)
            
            return Response(datos_turistas, status=status.HTTP_200_OK)

        except json.JSONDecodeError:
            return Response({"error": "La IA no devolvió un JSON válido.", "raw_response": response.text}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


# Colores suaves para distinguir grupos en el Excel (se repiten si hay muchos)
COLORES_GRUPO = ['C6EFCE', 'BDD7EE', 'FFEB9C', 'F8CBAD', 'D9D2E9', 'B4E5E6', 'FCE4D6', 'E2EFDA', 'DDEBF7', 'EAD1DC']


def exportar_excel(request, viaje_id, tipo):
    try:
        viaje = Viaje.objects.get(id=viaje_id)
    except Viaje.DoesNotExist:
        return HttpResponse("Viaje no encontrado", status=404)

    if tipo == 'logistica':
        # Juntos los de un mismo grupo; los que no tienen grupo al final
        turistas = Turista.objects.filter(viaje=viaje).order_by(F('grupo_id').asc(nulls_last=True), 'id')
    else:
        turistas = Turista.objects.filter(viaje=viaje).order_by('id')

    wb = Workbook()
    ws = wb.active
    ws.title = f"Lista_{tipo.capitalize()}"

    fuente_fecha = Font(size=36, bold=True, color="4F81BD", name="Calibri")
    fuente_jardin = Font(size=40, bold=True, color="00B050", name="Impact")

    fuente_cabecera = Font(bold=True, size=12, color="000000")
    fuente_normal = Font(size=11, color="000000")

    alineacion_centro = Alignment(horizontal='center', vertical='center')
    alineacion_izquierda = Alignment(horizontal='left', vertical='center')
    borde_delgado = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    ultima_columna = 'D' if tipo == 'cobro' else 'I'

    meses = {1: 'ENERO', 2: 'FEBRERO', 3: 'MARZO', 4: 'ABRIL', 5: 'MAYO', 6: 'JUNIO',
             7: 'JULIO', 8: 'AGOSTO', 9: 'SEPTIEMBRE', 10: 'OCTUBRE', 11: 'NOVIEMBRE', 12: 'DICIEMBRE'}

    ws.merge_cells(f'A1:{ultima_columna}1')
    celda_fecha = ws['A1']
    celda_fecha.value = f"{viaje.fecha_salida.day} DE {meses[viaje.fecha_salida.month]}"
    celda_fecha.font = fuente_fecha
    celda_fecha.alignment = alineacion_centro
    ws.row_dimensions[1].height = 45

    ws.merge_cells(f'A2:{ultima_columna}2')
    celda_jardin = ws['A2']
    celda_jardin.value = "JARDÍN DE LAS DELICIAS"
    celda_jardin.font = fuente_jardin
    celda_jardin.alignment = alineacion_centro
    ws.row_dimensions[2].height = 50

    ws.append([])

    colores_fila = {}  # fila -> color del grupo
    fila_totales = None

    if tipo == 'cobro':
        cabeceras = ['N', 'NOMBRES', 'EDAD', 'C.I.']
        ws.append(cabeceras)
        for i, t in enumerate(turistas, start=1):
            ws.append([i, t.nombre_completo, t.edad if t.edad else '', t.ci if t.ci else ''])

    elif tipo == 'logistica':
        cabeceras = ['N', 'NOMBRES', 'GRUPO', 'RESERVAS', 'DESC.', 'SALDO', 'RAPPEL', 'TIROLESA', 'OBS']
        ws.append(cabeceras)
        numero_grupo = {}
        for i, t in enumerate(turistas, start=1):
            if t.grupo_id and t.grupo_id not in numero_grupo:
                numero_grupo[t.grupo_id] = len(numero_grupo) + 1
            n = numero_grupo.get(t.grupo_id)
            rappel = 'SI' if t.hace_rappel else 'NO'
            tirolesa = 'SI' if t.hace_tirolesa else 'NO'
            ws.append([i, t.nombre_completo, f'G{n}' if n else '-', t.reserva, t.descuento or '',
                       t.saldo, rappel, tirolesa, t.observaciones])
            if n:
                colores_fila[ws.max_row] = COLORES_GRUPO[(n - 1) % len(COLORES_GRUPO)]

        lista = list(turistas)
        ws.append([
            '', 'TOTAL', f'{len(lista)} pax',
            sum(t.reserva for t in lista), sum(t.descuento for t in lista), sum(t.saldo for t in lista),
            sum(t.hace_rappel for t in lista), sum(t.hace_tirolesa for t in lista), '',
        ])
        fila_totales = ws.max_row

    for fila in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=1, max_col=len(cabeceras)):
        for celda in fila:
            celda.border = borde_delgado
            if celda.row == 4 or celda.row == fila_totales:
                celda.font = fuente_cabecera
                celda.alignment = alineacion_centro
            else:
                celda.font = fuente_normal
                if celda.column != 2:
                    celda.alignment = alineacion_centro
                else:
                    celda.alignment = alineacion_izquierda
            if celda.row in colores_fila:
                celda.fill = PatternFill(start_color=colores_fila[celda.row], end_color=colores_fila[celda.row], fill_type='solid')

    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 35
    if tipo == 'cobro':
        ws.column_dimensions['C'].width = 10
        ws.column_dimensions['D'].width = 20
    else:
        ws.column_dimensions['C'].width = 10
        for col in ['D', 'E', 'F', 'G', 'H']:
            ws.column_dimensions[col].width = 12
        ws.column_dimensions['I'].width = 25

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Lista_{tipo}_{viaje.fecha_salida}.xlsx"'
    wb.save(response)

    return response
